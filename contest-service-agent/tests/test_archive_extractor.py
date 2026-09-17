import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from archive_extractor import build_archive_catalog, sha256  # noqa: E402


def make_source_deck(path: Path, title: str) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for text in (title, "Authorized source text"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(4))
        box.text = text
    prs.save(path)


class ArchiveExtractorTests(unittest.TestCase):
    def test_directory_catalog_reports_duplicates_without_lyrics(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            make_source_deck(root / "one.pptx", "Same Song")
            make_source_deck(root / "two.pptx", "Same Song")
            approved = {name: sha256(root / name) for name in ("one.pptx", "two.pptx")}
            catalog = build_archive_catalog(root, approved)
            self.assertEqual(catalog["schema_version"], 2)
            self.assertEqual(catalog["deck_count"], 2)
            self.assertEqual(catalog["song_count"], 2)
            self.assertEqual(len(catalog["duplicates"]), 1)
            self.assertNotIn("Authorized source text", str(catalog))
            self.assertNotIn("Same Song", str(catalog))
            self.assertRegex(catalog["songs"][0]["title_fingerprint"], r"^[a-f0-9]{64}$")

    def test_unconfirmed_and_changed_decks_are_not_indexed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            make_source_deck(root / "approved.pptx", "Approved Song")
            make_source_deck(root / "unknown.pptx", "Unknown Song")
            wrong_hash = "0" * 64
            catalog = build_archive_catalog(root, {"approved.pptx": wrong_hash})
            self.assertEqual(catalog["deck_count"], 0)
            self.assertEqual(catalog["song_count"], 0)
            reasons = {item["source_deck"]: item["reason"] for item in catalog["unconfirmed_decks"]}
            self.assertEqual(reasons["approved.pptx"], "hash_mismatch")
            self.assertEqual(reasons["unknown.pptx"], "not_in_manifest")


if __name__ == "__main__":
    unittest.main()
