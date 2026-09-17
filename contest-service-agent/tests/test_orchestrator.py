import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from orchestrator import prepare_service  # noqa: E402
from service_agent import build_editable_song_deck, quality_check_deck, sha256_file  # noqa: E402


class OrchestratorTests(unittest.TestCase):
    def _catalog(self, root: Path) -> Path:
        catalog = root / "catalog.json"
        catalog.write_text(json.dumps({
            "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
            "segments": [],
        }), encoding="utf-8")
        return catalog

    def test_existing_missing_and_created_flow(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            existing_path = Path(build_editable_song_deck(
                "Existing Song", "User supplied text", str(archive / "Existing Song.pptx")
            ))
            catalog = root / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "segments": [{
                    "title": "Existing Song",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": existing_path.name,
                    "source_sha256": sha256_file(existing_path),
                    "start_slide": 1,
                    "end_slide": 1,
                }],
            }), encoding="utf-8")

            result = prepare_service(
                "Sunday Front Porch",
                [
                    {"title": "Existing Song"},
                    {"title": "Missing Song"},
                    {"title": "New Supplied Song", "lyrics": "Line one\n\nLine two"},
                ],
                str(catalog),
                str(archive),
                str(root / "out"),
            )

            self.assertEqual(result["status"], "AWAITING_PASTOR_APPROVAL")
            self.assertFalse(result["ready_for_approval"])
            self.assertEqual(result["items"][0]["status"], "REUSED_EXISTING")
            self.assertEqual(result["items"][1]["status"], "NEEDS_HUMAN_INPUT")
            self.assertEqual(result["items"][2]["status"], "CREATED_DRAFT")
            self.assertTrue(result["items"][2]["qa"]["ok"])
            self.assertTrue(Path(result["final_deck"]).is_file())
            self.assertEqual(result["summary"]["assembled_slides"], 3)

    def test_filename_match_cannot_bypass_confirmation_catalog(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            build_editable_song_deck("Unconfirmed Song", "Archive text", str(archive / "Unconfirmed Song.pptx"))
            result = prepare_service(
                "Sunday Front Porch",
                [{"title": "Unconfirmed Song"}],
                str(self._catalog(root)),
                str(archive),
                str(root / "out"),
            )
            self.assertEqual(result["items"][0]["status"], "NEEDS_HUMAN_INPUT")
            self.assertEqual(result["summary"]["assembled_slides"], 0)

    def test_changed_confirmed_deck_is_blocked(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            deck = Path(build_editable_song_deck("Confirmed Song", "Old text", str(archive / "Confirmed Song.pptx")))
            confirmed_hash = sha256_file(deck)
            build_editable_song_deck("Confirmed Song", "Changed text", str(deck))
            catalog = root / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "segments": [{
                    "title": "Confirmed Song",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": deck.name,
                    "source_sha256": confirmed_hash,
                    "start_slide": 1,
                    "end_slide": 1,
                }],
            }), encoding="utf-8")
            result = prepare_service(
                "Sunday Front Porch",
                [{"title": "Confirmed Song"}],
                str(catalog),
                str(archive),
                str(root / "out"),
            )
            self.assertEqual(result["items"][0]["status"], "SOURCE_CONFIRMATION_FAILED")

    def test_confirmed_source_is_copied_from_one_hashed_snapshot(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            deck = Path(build_editable_song_deck("Confirmed Song", "Approved text", str(archive / "Confirmed Song.pptx")))
            catalog = root / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "segments": [{
                    "title": "Confirmed Song",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": deck.name,
                    "source_sha256": sha256_file(deck),
                    "start_slide": 1,
                    "end_slide": 1,
                }],
            }), encoding="utf-8")
            changed = False

            def mutate_source_after_snapshot(path, *args, **kwargs):
                nonlocal changed
                if "source-snapshots" in str(path) and not changed:
                    changed = True
                    build_editable_song_deck("Confirmed Song", "Replacement text", str(deck))
                return quality_check_deck(path, *args, **kwargs)

            with patch("orchestrator.quality_check_deck", side_effect=mutate_source_after_snapshot):
                result = prepare_service(
                    "Sunday Front Porch",
                    [{"type": "song", "title": "Confirmed Song"}],
                    str(catalog),
                    str(archive),
                    str(root / "out"),
                )
            self.assertTrue(result["ready_for_approval"], result)
            assembled = Presentation(result["final_deck"])
            assembled_text = "\n".join(shape.text for slide in assembled.slides for shape in slide.shapes if hasattr(shape, "text"))
            self.assertIn("Approved text", assembled_text)
            self.assertNotIn("Replacement text", assembled_text)
            self.assertEqual(list((Path(result["output_directory"]) / "source-snapshots").iterdir()), [])

    def test_confirmed_catalog_without_reviewed_range_cannot_copy_whole_deck(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            deck = Path(build_editable_song_deck(
                "Approved Song",
                "Approved song text\n\nPRIVATE unrelated slide",
                str(archive / "mixed-service.pptx"),
            ))
            catalog = root / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "segments": [{
                    "title": "Approved Song",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": deck.name,
                    "source_sha256": sha256_file(deck),
                }],
            }), encoding="utf-8")

            result = prepare_service(
                "Sunday Front Porch",
                [{"type": "song", "title": "Approved Song"}],
                str(catalog),
                str(archive),
                str(root / "out"),
            )

            self.assertEqual(result["items"][0]["status"], "NEEDS_HUMAN_INPUT")
            self.assertEqual(result["summary"]["assembled_slides"], 0)
            self.assertIsNone(result["final_deck"])

    def test_unselected_text_cannot_mask_image_only_song_range(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            image_path = root / "slide.png"
            Image.new("RGB", (80, 80), "black").save(image_path)
            source = Presentation()
            source.slide_width = Inches(13.333)
            source.slide_height = Inches(7.5)
            image_slide = source.slides.add_slide(source.slide_layouts[6])
            image_slide.shapes.add_picture(str(image_path), 0, 0)
            text_slide = source.slides.add_slide(source.slide_layouts[6])
            text_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text = "Unselected text"
            deck = archive / "mixed-service.pptx"
            source.save(deck)
            catalog = root / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "segments": [{
                    "title": "Image Only Song",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": deck.name,
                    "source_sha256": sha256_file(deck),
                    "start_slide": 1,
                    "end_slide": 1,
                }],
            }), encoding="utf-8")

            result = prepare_service(
                "Sunday Front Porch",
                [
                    {"type": "service_title", "title": "Welcome"},
                    {"type": "song", "title": "Image Only Song"},
                ],
                str(catalog),
                str(archive),
                str(root / "out"),
            )

            self.assertEqual(result["items"][1]["status"], "QA_FAILED")
            self.assertIn("no_editable_text", result["items"][1]["qa"]["errors"])
            self.assertFalse(result["ready_for_approval"])

    def test_complete_flow_stops_at_human_approval(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            result = prepare_service(
                "Sunday Front Porch",
                [{"title": "Authorized Draft", "lyrics": "Line one\nLine two"}],
                str(self._catalog(root)),
                str(archive),
                str(root / "out"),
            )
            self.assertTrue(result["ready_for_approval"], result)
            self.assertFalse(result["approved"])
            self.assertFalse(result["autopublish"])
            manifest = json.loads(Path(result["approval_manifest"]).read_text(encoding="utf-8"))
            self.assertEqual(manifest["schema_version"], 3)
            self.assertEqual(manifest["items"][0]["item_type"], "song")
            self.assertNotIn("songs", manifest)
            self.assertIsNone(manifest["approval"]["decision"])
            self.assertFalse(manifest["autopublish"])
            self.assertEqual(manifest["final_deck_sha256"], result["final_deck_sha256"])
            self.assertEqual(len(manifest["final_deck_sha256"]), 64)
            page = Path(result["approval_ui"]).read_text(encoding="utf-8")
            self.assertIn("cannot publish or send anything", page)
            self.assertIn(result["final_deck_sha256"], page)
            self.assertIn("APPROVAL_RECOMMENDED", page)
            self.assertIn("identity_verified:false", page)

    def test_ordered_front_porch_service_builds_all_supported_item_types(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            result = prepare_service(
                "Complete Front Porch",
                [
                    {"type": "service_title", "title": "Welcome to KCMC"},
                    {"type": "song", "title": "Authorized Song", "lyrics": "Verse\n\nChorus"},
                    {"type": "scripture", "title": "John 1:1", "text": "Authorized Scripture text"},
                    {"type": "sermon_title", "title": "The Sermon"},
                    {"type": "announcement", "title": "Community Meal", "text": "Friday at six"},
                    {"type": "blank", "title": "Camera transition"},
                ],
                str(self._catalog(root)),
                str(archive),
                str(root / "out"),
            )
            self.assertTrue(result["ready_for_approval"], result)
            self.assertEqual(
                [item["item_type"] for item in result["items"]],
                ["service_title", "song", "scripture", "sermon_title", "announcement", "blank"],
            )
            self.assertEqual(result["summary"]["assembled_slides"], 9)
            self.assertEqual(Path(result["output_directory"]).stat().st_mode & 0o777, 0o700)
            self.assertEqual(Path(result["final_deck"]).stat().st_mode & 0o777, 0o600)

    def test_each_build_uses_a_fresh_run_directory(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            output_root = root / "out"
            first = prepare_service(
                "Same Service",
                [{"type": "song", "title": "Ready", "lyrics": "Authorized line"}],
                str(self._catalog(root)),
                str(archive),
                str(output_root),
            )
            second = prepare_service(
                "Same Service",
                [{"type": "song", "title": "Blocked"}],
                str(self._catalog(root)),
                str(archive),
                str(output_root),
            )
            self.assertNotEqual(first["run_id"], second["run_id"])
            self.assertNotEqual(first["output_directory"], second["output_directory"])
            self.assertTrue(Path(first["final_deck"]).is_file())
            self.assertIsNone(second["final_deck"])
            second_page = Path(second["approval_ui"]).read_text(encoding="utf-8")
            self.assertNotIn("Record approval recommendation", second_page)

    def test_unverified_service_style_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            with self.assertRaises(ValueError):
                prepare_service(
                    "Unverified Style",
                    [{"type": "service_title", "title": "Title"}],
                    str(self._catalog(root)),
                    str(archive),
                    str(root / "out"),
                    service_style="Contemporary",
                )

    def test_output_inside_archive_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            with self.assertRaises(ValueError):
                prepare_service(
                    "Sunday Front Porch",
                    [{"title": "Song", "lyrics": "Line"}],
                    str(self._catalog(root)),
                    str(archive),
                    str(archive / "output"),
                )

    def test_nonexistent_archive_is_rejected_before_output_creation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "not-created" / "archive"
            output = archive / "output"
            with self.assertRaises(ValueError):
                prepare_service(
                    "Sunday Front Porch",
                    [{"title": "Song", "lyrics": "Line"}],
                    str(self._catalog(root)),
                    str(archive),
                    str(output),
                )
            self.assertFalse(output.exists())

    def test_malformed_song_item_is_blocked_without_crashing(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            result = prepare_service(
                "Sunday Front Porch",
                ["not an object"],
                str(self._catalog(root)),
                str(archive),
                str(root / "out"),
            )
            self.assertFalse(result["ready_for_approval"])
            self.assertEqual(result["items"][0]["status"], "NEEDS_HUMAN_INPUT")


if __name__ == "__main__":
    unittest.main()
