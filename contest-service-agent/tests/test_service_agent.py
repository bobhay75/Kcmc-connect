import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from service_agent import (  # noqa: E402
    append_blank_slide,
    append_slides_from_deck,
    build_editable_song_deck,
    find_song_in_catalog,
    normalize,
    quality_check_deck,
    resolve_archive_source,
)


class ServiceAgentTests(unittest.TestCase):
    def test_normalize(self):
        self.assertEqual(normalize("Great Are You, Lord!"), "greatareyoulord")

    def test_build_and_qa_black_white_editable_deck(self):
        with tempfile.TemporaryDirectory() as temp:
            deck = build_editable_song_deck("Test", "Verse one\nVerse two\n\nChorus", str(Path(temp) / "test.pptx"))
            qa = quality_check_deck(deck)
            self.assertTrue(qa["ok"], qa)
            self.assertEqual(qa["slides"], 2)
            self.assertGreaterEqual(qa["editable_text_shapes"], 2)
            prs = Presentation(deck)
            self.assertEqual(str(prs.slides[0].background.fill.fore_color.rgb), "000000")

    def test_blank_service_slide_is_black_and_structurally_drawable(self):
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        self.assertEqual(append_blank_slide(prs), 1)
        slide = prs.slides[0]
        self.assertEqual(str(slide.background.fill.fore_color.rgb), "000000")
        self.assertEqual(len(slide.shapes), 1)

    def test_catalog_supports_reviewed_songs_schema(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            catalog.write_text(json.dumps({
                "source_deck": "service.pptx",
                "source_sha256": "a" * 64,
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "songs": [{
                    "title": "Welcome Table",
                    "role": "song",
                    "service_style": "Front Porch",
                    "start_slide": 5,
                    "end_slide": 13,
                }],
            }), encoding="utf-8")
            result = find_song_in_catalog("Welcome Table", str(catalog))
            self.assertEqual(result["status"], "found")
            self.assertEqual(result["source"], "service.pptx")
            self.assertEqual(result["start_slide"], 5)

    def test_catalog_supports_private_title_fingerprints(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            fingerprint = hashlib.sha256(b"welcometable").hexdigest()
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "songs": [{
                    "title_fingerprint": fingerprint,
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": "service.pptx",
                    "source_sha256": "a" * 64,
                    "start_slide": 5,
                    "end_slide": 13,
                }],
            }), encoding="utf-8")
            result = find_song_in_catalog("Welcome Table", str(catalog))
            self.assertEqual(result["status"], "found")
            self.assertEqual(result["source"], "service.pptx")

    def test_unconfirmed_catalog_cannot_supply_archive_slides(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            catalog.write_text(json.dumps({
                "songs": [{"title": "Unverified Song", "source_deck": "unknown.pptx"}],
            }), encoding="utf-8")
            result = find_song_in_catalog("Unverified Song", str(catalog))
            self.assertEqual(result["status"], "unconfirmed_catalog")

    def test_slide_assembly(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = build_editable_song_deck("Amazing Grace", "Amazing grace", str(root / "Amazing Grace.pptx"))
            destination = Presentation()
            destination.slide_width = Presentation(source).slide_width
            destination.slide_height = Presentation(source).slide_height
            copied = append_slides_from_deck(destination, source)
            self.assertTrue(copied["ok"], copied)
            assembled = root / "assembled.pptx"
            destination.save(assembled)
            self.assertTrue(quality_check_deck(str(assembled))["ok"])

    def test_assembly_keeps_distinct_images_from_multiple_decks(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            destination = Presentation()
            destination.slide_width = Inches(13.333)
            destination.slide_height = Inches(7.5)
            expected_hashes = []

            for color in ("red", "blue"):
                image_path = root / f"{color}.png"
                Image.new("RGB", (80, 80), color).save(image_path)
                source = Presentation()
                source.slide_width = destination.slide_width
                source.slide_height = destination.slide_height
                slide = source.slides.add_slide(source.slide_layouts[6])
                picture = slide.shapes.add_picture(str(image_path), 0, 0)
                expected_hashes.append(picture.image.sha1)
                source_path = root / f"{color}.pptx"
                source.save(source_path)
                copied = append_slides_from_deck(destination, str(source_path))
                self.assertTrue(copied["ok"], copied)

            assembled = root / "images.pptx"
            destination.save(assembled)
            reopened = Presentation(assembled)
            actual_hashes = [
                shape.image.sha1
                for slide in reopened.slides
                for shape in slide.shapes
                if getattr(shape, "image", None) is not None
            ]
            self.assertEqual(actual_hashes, expected_hashes)

    def test_external_relationship_deck_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = Presentation()
            source.slide_width = Inches(13.333)
            source.slide_height = Inches(7.5)
            slide = source.slides.add_slide(source.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
            box.text = "Linked text"
            box.click_action.hyperlink.address = "https://example.invalid/tracker"
            source_path = root / "external-link.pptx"
            source.save(source_path)

            qa = quality_check_deck(str(source_path), expected_font=None, expected_font_size_pt=None)
            self.assertIn("unsafe_external_relationship", qa["errors"])
            destination = Presentation()
            destination.slide_width = source.slide_width
            destination.slide_height = source.slide_height
            copied = append_slides_from_deck(destination, str(source_path))
            self.assertFalse(copied["ok"])
            self.assertEqual(copied["error"], "unsafe_powerpoint_relationships")

    def test_catalog_path_cannot_escape_archive(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            archive = root / "archive"
            archive.mkdir()
            outside = build_editable_song_deck("Outside", "Text", str(root / "outside.pptx"))
            result = resolve_archive_source(outside, str(archive))
            self.assertEqual(result["status"], "blocked_path")

    def test_invalid_catalog_slide_range_is_reported(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "songs": [{
                    "title": "Bad Range",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": "deck.pptx",
                    "source_sha256": "a" * 64,
                    "start_slide": "oops",
                    "end_slide": 4,
                }],
            }), encoding="utf-8")
            result = find_song_in_catalog("Bad Range", str(catalog))
            self.assertEqual(result["status"], "invalid_catalog")

    def test_catalog_range_rejects_boolean_and_fractional_boundaries(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            for invalid_start in (True, 1.9):
                with self.subTest(invalid_start=invalid_start):
                    catalog.write_text(json.dumps({
                        "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                        "songs": [{
                            "title": "Exact Range",
                            "role": "song",
                            "service_style": "Front Porch",
                            "source_deck": "deck.pptx",
                            "source_sha256": "a" * 64,
                            "start_slide": invalid_start,
                            "end_slide": 2,
                        }],
                    }), encoding="utf-8")
                    result = find_song_in_catalog("Exact Range", str(catalog))
                    self.assertEqual(result["status"], "invalid_catalog")

    def test_catalog_entry_without_reviewed_range_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-and-human-reviewed-ranges"},
                "songs": [{
                    "title": "Unsafe Whole Deck",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": "service.pptx",
                    "source_sha256": "a" * 64,
                }],
            }), encoding="utf-8")
            result = find_song_in_catalog("Unsafe Whole Deck", str(catalog))
            self.assertEqual(result["status"], "invalid_catalog")

    def test_legacy_hash_only_confirmation_method_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            catalog = Path(temp) / "catalog.json"
            catalog.write_text(json.dumps({
                "confirmation": {"status": "confirmed", "method": "sha256-manifest"},
                "songs": [{
                    "title": "Legacy Entry",
                    "role": "song",
                    "service_style": "Front Porch",
                    "source_deck": "service.pptx",
                    "source_sha256": "a" * 64,
                    "start_slide": 1,
                    "end_slide": 1,
                }],
            }), encoding="utf-8")
            result = find_song_in_catalog("Legacy Entry", str(catalog))
            self.assertEqual(result["status"], "unconfirmed_catalog")


if __name__ == "__main__":
    unittest.main()
