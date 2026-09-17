from __future__ import annotations

import json
import hashlib
import re
import time
from datetime import datetime, timezone
from dataclasses import asdict, dataclass
from pathlib import Path
from uuid import uuid4

from pptx import Presentation

from service_agent import (
    KCMC_SLIDE_HEIGHT,
    KCMC_SLIDE_WIDTH,
    SUPPORTED_ITEM_TYPES,
    SUPPORTED_SERVICE_STYLES,
    append_blank_slide,
    append_slides_from_deck,
    build_editable_song_deck,
    build_editable_text_deck,
    find_song_in_catalog,
    quality_check_deck,
    resolve_archive_source,
    sha256_file,
    write_approval_manifest,
    write_approval_ui,
)

MAX_APPROVED_DECK_BYTES = 100 * 1024 * 1024


@dataclass
class ServiceItem:
    title: str
    status: str
    item_type: str = "song"
    source: str | None = None
    slides_added: int = 0
    qa: dict | None = None
    needs_lyrics: bool = False
    needs_selection: bool = False
    detail: str | None = None


def safe_slug(value: str, fallback: str = "service") -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return (slug[:80].strip("-") or fallback)


def _source_label(path: str, archive_root: str) -> str:
    source = Path(path).resolve()
    root = Path(archive_root).resolve()
    if source.is_relative_to(root):
        return str(source.relative_to(root))
    return source.name


def _new_deck() -> Presentation:
    deck = Presentation()
    deck.slide_width = KCMC_SLIDE_WIDTH
    deck.slide_height = KCMC_SLIDE_HEIGHT
    return deck


def _append_existing(
    final_deck: Presentation,
    title: str,
    match: dict,
    archive_root: str,
    snapshot_dir: Path,
) -> ServiceItem | None:
    source_result = resolve_archive_source(match.get("source"), archive_root)
    if source_result.get("status") != "found":
        return None
    source_path = str(source_result["path"])
    expected_sha256 = str(match.get("source_sha256") or "").lower()
    if not re.fullmatch(r"[a-f0-9]{64}", expected_sha256):
        return ServiceItem(
            title=title,
            status="SOURCE_CONFIRMATION_FAILED",
            source=_source_label(source_path, archive_root),
            needs_selection=True,
            detail="Catalog match is missing its confirmed source hash.",
        )
    try:
        with Path(source_path).open("rb") as source_file:
            source_bytes = source_file.read(MAX_APPROVED_DECK_BYTES + 1)
    except OSError:
        source_bytes = b""
    if len(source_bytes) > MAX_APPROVED_DECK_BYTES:
        return ServiceItem(
            title=title,
            status="SOURCE_CONFIRMATION_FAILED",
            source=_source_label(source_path, archive_root),
            needs_selection=True,
            detail="Confirmed archive deck exceeds the 100 MB safety limit.",
        )
    if hashlib.sha256(source_bytes).hexdigest() != expected_sha256:
        return ServiceItem(
            title=title,
            status="SOURCE_CONFIRMATION_FAILED",
            source=_source_label(source_path, archive_root),
            needs_selection=True,
            detail="Archive deck changed after confirmation; re-confirm and rebuild the private catalog.",
        )
    snapshot_path = snapshot_dir / f"{expected_sha256}.pptx"
    snapshot = str(snapshot_path)
    try:
        snapshot_path.write_bytes(source_bytes)
        snapshot_path.chmod(0o600)
        source_qa = quality_check_deck(
            snapshot,
            expected_font=None,
            expected_font_size_pt=None,
            start_slide=int(match["start_slide"]),
            end_slide=int(match["end_slide"]),
        )
        if not source_qa.get("ok"):
            return ServiceItem(
                title=title,
                status="QA_FAILED",
                source=_source_label(source_path, archive_root),
                qa=source_qa,
                detail="The selected archive range failed editability or layout checks.",
            )
        copied = append_slides_from_deck(
            final_deck,
            snapshot,
            int(match["start_slide"]),
            int(match["end_slide"]),
        )
        if not copied.get("ok"):
            return ServiceItem(
                title=title,
                status="QA_FAILED",
                source=_source_label(source_path, archive_root),
                qa={"ok": False, "errors": [copied.get("error")], "warnings": []},
                detail="The selected archive slide range could not be assembled.",
            )
        return ServiceItem(
            title=title,
            status="REUSED_EXISTING",
            source=_source_label(source_path, archive_root),
            slides_added=int(copied["slides_added"]),
            qa=source_qa,
        )
    finally:
        snapshot_path.unlink(missing_ok=True)


def prepare_service(
    service_name: str,
    service_items: list[dict],
    catalog_path: str,
    archive_root: str,
    output_root: str,
    service_style: str = "Front Porch",
) -> dict:
    """Build one editable service deck and stop at the human approval gate.

    Ordered items may be service titles, songs, Scripture, sermon titles,
    announcements, or intentional blanks. It never fetches or invents text.
    """
    service_name = service_name.strip()
    if not service_name:
        raise ValueError("Service name is required.")
    if service_style not in SUPPORTED_SERVICE_STYLES:
        raise ValueError("Only the verified Front Porch style is currently supported.")
    if not isinstance(service_items, list) or not service_items:
        raise ValueError("At least one ordered service item is required.")
    started = time.monotonic()
    output_base = Path(output_root).expanduser().resolve()
    archive = Path(archive_root).expanduser().resolve()
    if not archive.is_dir():
        raise ValueError("Approved private archive directory does not exist.")
    if output_base == archive or output_base.is_relative_to(archive):
        raise ValueError("Output directory must be outside the private source archive.")
    if output_base.exists() and not output_base.is_dir():
        raise ValueError("Output root must be a directory.")
    output_base.mkdir(parents=True, exist_ok=True, mode=0o700)
    run_id = f"{datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')}-{uuid4().hex[:8]}"
    out = output_base / run_id
    out.mkdir(mode=0o700)
    asset_dir = out / "draft-assets"
    asset_dir.mkdir(mode=0o700)
    snapshot_dir = out / "source-snapshots"
    snapshot_dir.mkdir(mode=0o700)
    final_deck = _new_deck()
    items: list[ServiceItem] = []

    for position, item in enumerate(service_items, start=1):
        if not isinstance(item, dict):
            items.append(
                ServiceItem(
                    title=f"Invalid item {position}",
                    status="NEEDS_HUMAN_INPUT",
                    item_type="invalid",
                    detail="Each service item must be an object.",
                    needs_selection=True,
                )
            )
            continue
        item_type = str(item.get("type") or "song").strip().lower()
        if item_type not in SUPPORTED_ITEM_TYPES:
            items.append(
                ServiceItem(
                    title=str(item.get("title") or f"Item {position}"),
                    status="NEEDS_HUMAN_INPUT",
                    item_type=item_type or "invalid",
                    detail="Unsupported service item type.",
                    needs_selection=True,
                )
            )
            continue
        if item_type == "blank":
            append_blank_slide(final_deck)
            items.append(
                ServiceItem(
                    title=str(item.get("title") or "Intentional blank"),
                    status="CREATED_DRAFT",
                    item_type="blank",
                    slides_added=1,
                    qa={"ok": True, "errors": [], "warnings": []},
                )
            )
            continue

        title = str(item.get("title") or "").strip()
        if not title:
            items.append(
                ServiceItem(
                    title=f"Untitled item {position}",
                    status="NEEDS_HUMAN_INPUT",
                    item_type=item_type,
                    detail=f"Title is required for {item_type}.",
                    needs_selection=True,
                )
            )
            continue

        if item_type != "song":
            raw_text = item.get("text")
            text = raw_text.strip() if isinstance(raw_text, str) else ""
            if item_type in {"scripture", "announcement"} and not text:
                items.append(
                    ServiceItem(
                        title=title,
                        status="NEEDS_HUMAN_INPUT",
                        item_type=item_type,
                        detail=f"Authorized text is required for {item_type}.",
                    )
                )
                continue
            if item_type in {"scripture", "announcement"}:
                text = f"{title}\n\n{text}"
            elif text:
                text = f"{title}\n{text}"
            else:
                text = title
            name = f"{position:02d}-{item_type}-{safe_slug(title, item_type)}.pptx"
            deck_path = build_editable_text_deck(
                title,
                text,
                str(asset_dir / name),
                item_type,
                service_style,
            )
            qa = quality_check_deck(deck_path)
            copied = append_slides_from_deck(final_deck, deck_path) if qa.get("ok") else {"ok": False}
            status = "CREATED_DRAFT" if copied.get("ok") else "QA_FAILED"
            items.append(
                ServiceItem(
                    title=title,
                    status=status,
                    item_type=item_type,
                    source=str(Path("draft-assets") / name),
                    slides_added=int(copied.get("slides_added") or 0),
                    qa=qa,
                    detail=None if status == "CREATED_DRAFT" else "Generated content failed QA or assembly.",
                )
            )
            continue

        match = find_song_in_catalog(title, catalog_path)
        if match.get("status") == "ambiguous":
            items.append(
                ServiceItem(
                    title=title,
                    status="NEEDS_HUMAN_SELECTION",
                    item_type="song",
                    needs_selection=True,
                    detail="More than one catalog match has the same score.",
                )
            )
            continue

        existing = (
            _append_existing(final_deck, title, match, str(archive), snapshot_dir)
            if match.get("status") == "found"
            else None
        )

        if existing is not None:
            items.append(existing)
            continue

        raw_lyrics = item.get("lyrics")
        lyrics = raw_lyrics.strip() if isinstance(raw_lyrics, str) else ""
        if not lyrics:
            items.append(
                ServiceItem(
                    title=title,
                    status="NEEDS_HUMAN_INPUT",
                    item_type="song",
                    needs_lyrics=True,
                    detail="Authorized lyrics or an approved archive deck are required.",
                )
            )
            continue

        name = f"{position:02d}-{safe_slug(title, 'song')}.pptx"
        deck_path = build_editable_song_deck(title, lyrics, str(asset_dir / name), service_style)
        qa = quality_check_deck(deck_path)
        slides_added = 0
        status = "CREATED_DRAFT"
        detail = None
        if qa.get("ok"):
            copied = append_slides_from_deck(final_deck, deck_path)
            if copied.get("ok"):
                slides_added = int(copied["slides_added"])
            else:
                status = "QA_FAILED"
                detail = "Generated slides could not be added to the complete service deck."
                qa = {"ok": False, "errors": [copied.get("error")], "warnings": []}
        else:
            status = "QA_FAILED"
            detail = "Generated deck failed editability or layout checks."
        items.append(
            ServiceItem(
                title=title,
                status=status,
                item_type="song",
                source=str(Path("draft-assets") / name),
                slides_added=slides_added,
                qa=qa,
                detail=detail,
            )
        )

    payload = [asdict(item) for item in items]
    final_path: Path | None = None
    final_deck_sha256: str | None = None
    final_qa: dict | None = None
    if len(final_deck.slides):
        final_path = out / f"{safe_slug(service_name, 'kcmc-service')}.pptx"
        final_deck.core_properties.title = service_name
        final_deck.core_properties.subject = "KCMC approval draft; human approval required"
        final_deck.save(final_path)
        final_path.chmod(0o600)
        final_deck_sha256 = sha256_file(final_path)
        final_qa = quality_check_deck(str(final_path), expected_font=None, expected_font_size_pt=None)

    complete_statuses = {"REUSED_EXISTING", "CREATED_DRAFT"}
    ready_for_approval = bool(items) and all(
        item.status in complete_statuses and (item.qa is None or item.qa.get("ok")) for item in items
    ) and bool(final_qa and final_qa.get("ok"))

    manifest_path = out / "approval.json"
    manifest = write_approval_manifest(
        service_name,
        payload,
        str(manifest_path),
        final_path.name if final_path else None,
        final_deck_sha256,
        final_qa,
    )
    approval_ui = write_approval_ui(
        service_name,
        payload,
        str(out / "approval.html"),
        final_path.name if final_path else None,
        final_deck_sha256,
        ready_for_approval,
    )
    result = {
        "schema_version": 3,
        "run_id": run_id,
        "output_directory": str(out),
        "service": service_name,
        "service_style": service_style,
        "status": "AWAITING_PASTOR_APPROVAL",
        "items": payload,
        "summary": {
            "requested": len(items),
            "reused": sum(item.status == "REUSED_EXISTING" for item in items),
            "created": sum(item.status == "CREATED_DRAFT" for item in items),
            "blocked": sum(item.status not in complete_statuses for item in items),
            "assembled_slides": len(final_deck.slides),
        },
        "final_deck": str(final_path) if final_path else None,
        "final_deck_sha256": final_deck_sha256,
        "final_qa": final_qa,
        "approval_manifest": manifest,
        "approval_ui": approval_ui,
        "ready_for_approval": ready_for_approval,
        "approved": False,
        "autopublish": False,
        "pipeline_seconds": round(time.monotonic() - started, 3),
    }
    report_path = out / "production-report.json"
    report_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
    report_path.chmod(0o600)
    return result
